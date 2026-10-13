#!/usr/bin/env python3
# -*- coding: ascii -*-
"""
toolbox.py -- Treebeard Phase 2, Toolbox v1 (BRIDGE MODE)

The toolbox is the runtime for promoted tools: the ROOM they execute in and
the DOOR they reach the world through. The shelf (the toolkit itself, at
TOOLKIT_ROOT) already exists -- Anvil's lock puts approved code there. This
file adds the other two jobs.

DESIGN TRACE (operator spec, 2026-07-02/03)
  T1 SHELF (read-only here)
     -> Tools live at TOOLKIT_ROOT/<tool>/ with a main.py entry point.
        The toolbox never writes to the toolkit; only Anvil's lock does.
  T2 ROOM (the wall)
     -> Every run executes inside rootless Podman: --network=none,
        --userns=keep-id (container uid == host uid, so scratch comes back
        owned by the invoker and cleanup rmtree works -- the known Anvil /
        Meeseeks defect fixed by construction, not patched after),
        --cap-drop=ALL, --read-only rootfs, tool dir mounted :ro, a scratch
        dir :rw, resource caps, and a host-side kill on the wall clock.
        _room_cmd() is the single seam that builds the podman argv; nothing
        else in the file knows the container flags.
     -> Selftest uses an UNSAFE subprocess room (same lifecycle, no wall)
        so the broker, budget, ledger, and manifest logic are proven
        without Podman. Real runs REFUSE the unsafe room unless
        TOOLBOX_UNSAFE_ROOM=1 is set explicitly; the run row and status
        card then say UNSAFE in plain text. Never for real work.
  T3 DOOR (the audited hole)
     -> One Unix socket per run, bind-mounted into the container at
        /door/door.sock. Protocol: one JSON line request, one JSON line
        response, one request per connection.
          {"v":1, "skill":"router", "args":{...}}
          {"ok":true, "result":...} | {"ok":false, "error":"..."}
        The broker (host side) checks the tool's manifest allowlist and
        budget, SPENDS a ledger row before executing (attempts count --
        same discipline as the Meeseeks budget), executes host-side
        holding any keys, records the outcome, and passes text back in.
        Keys never enter the room. A skill not in the manifest is refused
        loud and audited; refusals are themselves capped per run.

  SKILLS (door services)
    S1 router     -> host-side call into TOOLBOX_LIB router.route() with a
                     manifest-listed tier. Frontier is reachable only if a
                     manifest says so; enroll grants nothing.
                     NOTE: result attribute names (ok/text/tier/model/
                     error) are read defensively; CONFIRM against the
                     deployed router.py result class in the walk before
                     the first live router-skill call.
    S2 state      -> returns the text of a registered status card (the
                     files Anvil / Meeseeks / this tool already write
                     atomically). Registry at ROOT/config/cards.json,
                     treebeard-owned. No arbitrary path reads -- the door
                     is not a filesystem primitive.
    S3 web_search -> single-provider egress adapter (Brave shape). SHIPS
                     DARK: no key on the box yet. CONFIRM endpoint + auth
                     header against current provider docs BEFORE
                     provisioning a key (same rule router.py carries for
                     Cheap_provider). TOOLBOX_SEARCH_FAKE=<json file> exercises
                     the plumbing without egress (selftest path).

  KNOWLEDGE SKILLS (not door services)
    pdf / xlsx / pptx are a venv library plus a doc; the model writes code
    using the library and the room runs it. init writes the docs to
    ROOT/skills/ (plus the tool-author contract). Callers (the Meeseeks
    box, agent wrappers) prepend the relevant doc to prompts. The venv
    packages and the room image that carries them are deploy-walk steps.

  GOVERNANCE
    -> Manifests are treebeard-owned JSON at ROOT/manifests/<tool>.json.
       No manifest, no run. enroll writes a deny-all manifest (empty
       skills, invokers=[gate user]). allow/revoke/invokers rewrite it
       atomically and audit. All manifest-editing commands require the
       gate user at the OS layer (getpass.getuser()), same reasoning as
       the anvil and meeseeks wrappers. Running a tool requires the
       caller to be in that tool's invokers list.
    -> Disciplines carried whole from anvil.py / meeseeks.py: single-
       transaction state+audit writes, temp-then-rename everywhere,
       staging provably empty, atomic status card at
       ROOT/status/toolbox.txt on every transition (Sun ingestion),
       fail-closed budgets, plain ASCII throughout.

BRIDGE MODE
  Subcommands are invoked manually from the treebeard shell. No daemon,
  no timer. Router-skill availability requires provider keys in the
  broker's environment at invocation (source treebeard.env), exactly as
  the Meeseeks live runs required yesterday.

ENV
  TOOLBOX_ROOT         data tree root       (default: <codedir>/data)
  TOOLBOX_LIB          router.py directory  (default: /opt/treebeard/shared/lib)
  TOOLKIT_ROOT         the shelf            (default: /opt/treebeard/shared/toolkit)
  TOOLBOX_IMAGE        room image           (default: localhost/treebeard-toolroom:latest)
  TOOLBOX_GATE_USER    gate identity        (default: treebeard)
  TOOLBOX_UNSAFE_ROOM  1 = subprocess room  (selftest / never for real)
  TOOLBOX_SEARCH_FAKE  canned search JSON   (selftest / offline)
  BRAVE_API_KEY        search key           (dark until provisioned)

COMMANDS
  init | enroll <tool> | allow <tool> <skill> [k=v ...] | revoke <tool> <skill>
  invokers <tool> <u1,u2> | run <tool> [args...] | status [tool]
  skills | cards | probe | selftest
"""

import getpass
import hashlib
import json
import os
import re
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import urllib.request

# ---------------------------------------------------------------------------
# constants and env
# ---------------------------------------------------------------------------

CODE_DIR = os.path.dirname(os.path.abspath(__file__))

# Door protocol hard caps. A request is one line; a response is one line.
REQ_LINE_CAP = 65536          # bytes, inbound request line
RESP_LINE_CAP = 1000000       # bytes, outbound response line
CARD_READ_CAP = 65536         # bytes, state-skill card read
LOG_SHOW_CAP = 4000           # bytes of stdout/stderr shown by `status`
REFUSED_CAP = 20              # refused door calls per run before hard-drop
DEFAULT_WALL = 120            # seconds, manifest may override
DEFAULT_DOOR_CAP = 16         # total spent door calls per run, manifest may override
ROUTER_MAX_TOKENS_CEIL = 4000 # absolute ceiling regardless of manifest
SEARCH_COUNT_CEIL = 8         # max results per search regardless of args

KNOWN_SKILLS = ("router", "state", "web_search")

_TOOL_RE = re.compile(r"^[a-z0-9][a-z0-9_-]{0,39}$")
_CARD_RE = re.compile(r"^[a-z0-9][a-z0-9_-]{0,39}$")

# Room resource caps (constants in v1; a later manifest key can override
# once a real tool needs it -- do not pre-build knobs nothing turns).
ROOM_PIDS = "64"
ROOM_MEM = "512m"
ROOM_CPUS = "1"
ROOM_TMPFS = "/tmp:rw,size=64m"


def ROOT():
    return os.environ.get("TOOLBOX_ROOT", os.path.join(CODE_DIR, "data"))


def LIB():
    return os.environ.get("TOOLBOX_LIB", "/opt/treebeard/shared/lib")


def TOOLKIT():
    return os.environ.get("TOOLKIT_ROOT", "/opt/treebeard/shared/toolkit")


def IMAGE():
    return os.environ.get("TOOLBOX_IMAGE", "localhost/treebeard-toolroom:latest")


def GATE_USER():
    return os.environ.get("TOOLBOX_GATE_USER", "treebeard")


def now():
    # UTC ISO, second resolution is enough for the ledger; the DB rowids
    # order same-second events.
    return time.strftime("%Y-%m-%dT%H:%M:%S+00:00", time.gmtime())


def _p(*parts):
    return os.path.join(ROOT(), *parts)


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def digest_obj(obj):
    # Stable digest of an args object for the ledger. Values are not
    # stored verbatim (a search query is fine; a router prompt is long),
    # only attributable.
    s = json.dumps(obj, sort_keys=True, ensure_ascii=True)
    return hashlib.sha256(s.encode("ascii", "replace")).hexdigest()[:16]


def atomic_write(path, text):
    # temp-then-rename, always. The temp lives in staging so "staging is
    # provably empty" covers torn writes too.
    os.makedirs(os.path.dirname(path), exist_ok=True)
    staging = _p("staging")
    os.makedirs(staging, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix="tb-", suffix=".tmp", dir=staging)
    os.fchmod(fd, 0o640)  # mkstemp creates 0600 regardless of umask; enforce tree convention
    try:
        with os.fdopen(fd, "w", encoding="ascii", errors="replace") as f:
            f.write(text)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def require_ascii(text, what):
    try:
        text.encode("ascii")
    except UnicodeEncodeError as e:
        raise ValueError("%s is not plain ASCII: %s" % (what, e))


# ---------------------------------------------------------------------------
# ledger
# ---------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS runs(
  rid TEXT PRIMARY KEY, tool TEXT NOT NULL, argv TEXT NOT NULL,
  state TEXT NOT NULL, room TEXT, image TEXT,
  created TEXT NOT NULL, started TEXT, ended TEXT,
  exit_code INTEGER, invoker TEXT, note TEXT);
CREATE TABLE IF NOT EXISTS door_calls(
  rid TEXT NOT NULL, seq INTEGER NOT NULL, skill TEXT NOT NULL,
  args_digest TEXT NOT NULL, state TEXT NOT NULL,
  started TEXT NOT NULL, ended TEXT, ok INTEGER, error TEXT,
  tier TEXT, model TEXT,
  PRIMARY KEY(rid, seq));
CREATE TABLE IF NOT EXISTS events(
  eid INTEGER PRIMARY KEY AUTOINCREMENT,
  ts TEXT NOT NULL, kind TEXT NOT NULL, rid TEXT, detail TEXT);
CREATE TABLE IF NOT EXISTS meta(k TEXT PRIMARY KEY, v TEXT);
"""


def db():
    # One short-lived connection per operation. Threads (the broker) open
    # their own; sqlite's default same-thread check then never fires.
    path = _p("ledger.db")
    conn = sqlite3.connect(path, timeout=10)
    conn.execute("PRAGMA busy_timeout=5000")
    return conn


def event(conn, kind, rid, detail):
    conn.execute(
        "INSERT INTO events(ts,kind,rid,detail) VALUES(?,?,?,?)",
        (now(), kind, rid, detail),
    )


def audited(kind, rid, detail, fn=None):
    # Single transaction: the state change (fn) and its audit row commit
    # together or not at all. The anvil transition() discipline.
    conn = db()
    try:
        with conn:
            if fn is not None:
                fn(conn)
            event(conn, kind, rid, detail)
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# status card
# ---------------------------------------------------------------------------

def write_card():
    conn = db()
    try:
        cur = conn.execute("SELECT COUNT(*) FROM runs")
        total = cur.fetchone()[0]
        cur = conn.execute(
            "SELECT COUNT(*) FROM runs WHERE state IN ('PENDING','RUNNING')")
        active = cur.fetchone()[0]
        cur = conn.execute(
            "SELECT rid,tool,state,room,created FROM runs "
            "ORDER BY created DESC, rid DESC LIMIT 5")
        recent = cur.fetchall()
    finally:
        conn.close()
    lines = ["TOOLBOX -- %s" % now()]
    lines.append("runs: total %d, active %d | room image: %s"
                 % (total, active, IMAGE()))
    unsafe = [r for r in recent if r[3] == "UNSAFE-SUBPROCESS"]
    if unsafe:
        lines.append("WARNING: recent runs used the UNSAFE room (no wall)")
    for rid, tool, state, room, created in recent:
        lines.append("  %s %s [%s] room=%s %s"
                     % (rid, tool, state, room or "-", created))
    try:
        cards = sorted(load_cards().keys())
        lines.append("cards: %s" % ",".join(cards))
    except Exception:
        lines.append("cards: (registry unreadable)")
    docs = []
    sk = _p("skills")
    if os.path.isdir(sk):
        docs = sorted(n[:-3] for n in os.listdir(sk) if n.endswith(".md"))
    lines.append("skills docs: %s" % ",".join(docs))
    atomic_write(_p("status", "toolbox.txt"), "\n".join(lines) + "\n")


# ---------------------------------------------------------------------------
# manifests (policy; treebeard-owned)
# ---------------------------------------------------------------------------

def gate(action):
    u = getpass.getuser()
    if u != GATE_USER():
        raise PermissionError(
            "%s requires the %s user at the OS layer (you are %s)"
            % (action, GATE_USER(), u))


def manifest_path(tool):
    return _p("manifests", "%s.json" % tool)


def load_manifest(tool):
    path = manifest_path(tool)
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="ascii") as f:
        m = json.load(f)
    validate_manifest(m, tool)
    return m


def validate_manifest(m, tool):
    if not isinstance(m, dict):
        raise ValueError("manifest for %s is not an object" % tool)
    if m.get("tool") != tool:
        raise ValueError("manifest tool field mismatch for %s" % tool)
    if not isinstance(m.get("invokers"), list) or not m["invokers"]:
        raise ValueError("manifest for %s needs a non-empty invokers list" % tool)
    if not isinstance(m.get("wall_seconds"), int) or m["wall_seconds"] < 1:
        raise ValueError("manifest for %s needs wall_seconds >= 1" % tool)
    if not isinstance(m.get("door_call_cap"), int) or m["door_call_cap"] < 0:
        raise ValueError("manifest for %s needs door_call_cap >= 0" % tool)
    skills = m.get("skills")
    if not isinstance(skills, dict):
        raise ValueError("manifest for %s needs a skills object" % tool)
    for name, cfg in skills.items():
        if name not in KNOWN_SKILLS:
            raise ValueError("manifest for %s grants unknown skill %r"
                             % (tool, name))
        if not isinstance(cfg, dict):
            raise ValueError("skill %s config must be an object" % name)


def save_manifest(m):
    validate_manifest(m, m["tool"])
    text = json.dumps(m, indent=2, sort_keys=True, ensure_ascii=True) + "\n"
    atomic_write(manifest_path(m["tool"]), text)


def cmd_enroll(tool):
    gate("enroll")
    check_tool_name(tool)
    tooldir = os.path.join(TOOLKIT(), tool)
    if not os.path.isdir(tooldir):
        raise FileNotFoundError("no such tool on the shelf: %s" % tooldir)
    if os.path.exists(manifest_path(tool)):
        raise FileExistsError("manifest already exists for %s" % tool)
    m = {
        "tool": tool,
        "invokers": [GATE_USER()],
        "wall_seconds": DEFAULT_WALL,
        "door_call_cap": DEFAULT_DOOR_CAP,
        "skills": {},  # deny-by-default: enroll grants NOTHING
    }
    save_manifest(m)
    audited("ENROLL", None, "tool=%s deny-all manifest written" % tool)
    write_card()
    print("enrolled %s: deny-all manifest at %s" % (tool, manifest_path(tool)))
    print("grant door services with: toolbox allow %s <skill> [k=v ...]" % tool)


def _parse_kv(pairs):
    out = {}
    for p in pairs:
        if "=" not in p:
            raise ValueError("expected key=value, got %r" % p)
        k, v = p.split("=", 1)
        out[k] = v
    return out


def cmd_allow(tool, skill, kvpairs):
    gate("allow")
    if skill not in KNOWN_SKILLS:
        raise ValueError("unknown skill %r (known: %s)"
                         % (skill, ",".join(KNOWN_SKILLS)))
    m = load_manifest(tool)
    if m is None:
        raise FileNotFoundError("no manifest for %s; enroll it first" % tool)
    kv = _parse_kv(kvpairs)
    cfg = {}
    if skill == "router":
        tiers = [t for t in kv.get("tiers", "").split(",") if t]
        if not tiers:
            raise ValueError("router grant needs tiers=a,b (explicit; "
                             "nothing is implied)")
        cfg["tiers"] = tiers
        cfg["max_calls"] = int(kv.get("max_calls", "4"))
        cfg["max_tokens"] = min(int(kv.get("max_tokens", "1200")),
                                ROUTER_MAX_TOKENS_CEIL)
    elif skill == "state":
        cards = [c for c in kv.get("cards", "").split(",") if c]
        if not cards:
            raise ValueError("state grant needs cards=a,b (explicit)")
        cfg["cards"] = cards
        cfg["max_calls"] = int(kv.get("max_calls", "8"))
    elif skill == "web_search":
        cfg["max_calls"] = int(kv.get("max_calls", "3"))
        cfg["max_results"] = min(int(kv.get("max_results", "5")),
                                 SEARCH_COUNT_CEIL)
    m["skills"][skill] = cfg
    save_manifest(m)
    audited("ALLOW", None, "tool=%s skill=%s cfg=%s"
            % (tool, skill, json.dumps(cfg, sort_keys=True)))
    print("allowed %s -> %s %s" % (tool, skill, json.dumps(cfg, sort_keys=True)))


def cmd_revoke(tool, skill):
    gate("revoke")
    m = load_manifest(tool)
    if m is None:
        raise FileNotFoundError("no manifest for %s" % tool)
    if skill not in m["skills"]:
        raise KeyError("%s does not have %s" % (tool, skill))
    del m["skills"][skill]
    save_manifest(m)
    audited("REVOKE", None, "tool=%s skill=%s" % (tool, skill))
    print("revoked %s -> %s" % (tool, skill))


def cmd_invokers(tool, users):
    gate("invokers")
    m = load_manifest(tool)
    if m is None:
        raise FileNotFoundError("no manifest for %s" % tool)
    lst = [u for u in users.split(",") if u]
    if not lst:
        raise ValueError("invokers needs at least one user")
    m["invokers"] = lst
    save_manifest(m)
    audited("INVOKERS", None, "tool=%s invokers=%s" % (tool, ",".join(lst)))
    print("invokers for %s: %s" % (tool, ",".join(lst)))


def check_tool_name(tool):
    if not _TOOL_RE.match(tool or ""):
        raise ValueError("bad tool name %r (want %s)" % (tool, _TOOL_RE.pattern))


# ---------------------------------------------------------------------------
# cards registry (state skill)
# ---------------------------------------------------------------------------

def cards_path():
    return _p("config", "cards.json")


def load_cards():
    with open(cards_path(), "r", encoding="ascii") as f:
        reg = json.load(f)
    if not isinstance(reg, dict):
        raise ValueError("cards.json must map name -> path")
    for name, path in reg.items():
        if not _CARD_RE.match(name):
            raise ValueError("bad card name %r in cards.json" % name)
        if not isinstance(path, str):
            raise ValueError("card %s path must be a string" % name)
    return reg


def seed_cards():
    reg = {
        "toolbox": _p("status", "toolbox.txt"),
        "meeseeks": "/opt/treebeard/shared/meeseeks/data/status/meeseeks.txt",
        "anvil": "/opt/treebeard/shared/anvil/data/status/ready.txt",
    }
    atomic_write(cards_path(),
                 json.dumps(reg, indent=2, sort_keys=True) + "\n")


# ---------------------------------------------------------------------------
# skills (broker side; the host performs these, never the room)
# ---------------------------------------------------------------------------

class SkillError(Exception):
    pass


def _router_mod():
    libp = LIB()
    if libp and libp not in sys.path:
        sys.path.insert(0, libp)
    try:
        import router  # noqa
        return router
    except Exception as e:
        raise SkillError("router library unavailable: %s" % e)


def skill_router(cfg, args):
    system = args.get("system", "")
    user = args.get("user", "")
    tier = args.get("tier", "")
    if not isinstance(user, str) or not user:
        raise SkillError("router: args.user (str) is required")
    if not isinstance(system, str):
        raise SkillError("router: args.system must be a string")
    if tier not in cfg.get("tiers", []):
        raise SkillError("router: tier %r not granted (manifest tiers: %s)"
                         % (tier, ",".join(cfg.get("tiers", []))))
    max_tokens = args.get("max_tokens", cfg.get("max_tokens", 800))
    try:
        max_tokens = min(int(max_tokens), int(cfg.get("max_tokens", 800)),
                         ROUTER_MAX_TOKENS_CEIL)
    except (TypeError, ValueError):
        raise SkillError("router: bad max_tokens")
    router = _router_mod()
    try:
        r = router.route(system, user, force_tier=tier)
    except Exception as e:
        raise SkillError("router raised: %s" % e)
    ok = bool(getattr(r, "ok", False))
    if not ok:
        reason = getattr(r, "error", None) or getattr(r, "reason", "") or "not-ok"
        raise SkillError("router not-ok: %s" % reason)
    # Defensive attribute read (v1): CONFIRM these names against the
    # deployed router result class in the walk. max_tokens is enforced
    # here by truncation as a belt; the router's own cap is the suspenders.
    text = getattr(r, "text", None)
    if text is None:
        text = getattr(r, "output", "")
    text = str(text)[: max_tokens * 8]  # crude char belt over token cap
    meta = {
        "tier": str(getattr(r, "tier", tier)),
        "model": str(getattr(r, "model", "") or ""),
    }
    return {"text": text, "tier": meta["tier"], "model": meta["model"]}, meta


def skill_state(cfg, args):
    name = args.get("card", "")
    if not _CARD_RE.match(name or ""):
        raise SkillError("state: bad card name")
    allowed = cfg.get("cards", [])
    if name not in allowed:
        raise SkillError("state: card %r not granted (manifest cards: %s)"
                         % (name, ",".join(allowed)))
    reg = load_cards()
    if name not in reg:
        raise SkillError("state: card %r not in registry" % name)
    path = reg[name]
    try:
        with open(path, "r", encoding="ascii", errors="replace") as f:
            text = f.read(CARD_READ_CAP)
        mtime = now() if not os.path.exists(path) else time.strftime(
            "%Y-%m-%dT%H:%M:%S+00:00", time.gmtime(os.path.getmtime(path)))
    except OSError as e:
        raise SkillError("state: cannot read card %s: %s" % (name, e))
    return {"card": name, "text": text, "mtime": mtime}, {}


# Search provider adapter. Brave shape; SHIPS DARK. CONFIRM the endpoint
# and auth header against current provider docs before provisioning a key
# -- provider strings drift (router.py carries the same warning for
# Cheap_provider). One provider, one key, every query a ledger row.
SEARCH_URL = "https://api.search.brave.com/res/v1/web/search"
SEARCH_KEY_ENVS = ("TOOLBOX_SEARCH_API_KEY", "BRAVE_API_KEY")


def skill_web_search(cfg, args):
    query = args.get("query", "")
    if not isinstance(query, str) or not query.strip():
        raise SkillError("web_search: args.query (str) is required")
    count = args.get("count", cfg.get("max_results", 5))
    try:
        count = min(int(count), int(cfg.get("max_results", 5)),
                    SEARCH_COUNT_CEIL)
    except (TypeError, ValueError):
        raise SkillError("web_search: bad count")
    fake = os.environ.get("TOOLBOX_SEARCH_FAKE")
    if fake:
        with open(fake, "r", encoding="ascii") as f:
            data = json.load(f)
        results = data.get("results", [])[:count]
        return {"query": query, "results": results, "provider": "fake"}, {}
    key = ""
    for envn in SEARCH_KEY_ENVS:
        key = os.environ.get(envn, "")
        if key:
            break
    if not key:
        raise SkillError("web_search unavailable: no key in %s (skill ships "
                         "dark until a provider key is provisioned)"
                         % "/".join(SEARCH_KEY_ENVS))
    url = SEARCH_URL + "?" + urllib.parse.urlencode(
        {"q": query, "count": str(count)})
    req = urllib.request.Request(url, headers={
        "Accept": "application/json",
        "X-Subscription-Token": key,
    })
    try:
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = json.loads(resp.read().decode("utf-8", "replace"))
    except Exception as e:
        raise SkillError("web_search failed: %s" % e)
    results = []
    for item in (data.get("web", {}) or {}).get("results", [])[:count]:
        results.append({
            "title": str(item.get("title", ""))[:300],
            "url": str(item.get("url", ""))[:500],
            "description": str(item.get("description", ""))[:500],
        })
    return {"query": query, "results": results, "provider": "brave"}, {}


SKILL_FNS = {
    "router": skill_router,
    "state": skill_state,
    "web_search": skill_web_search,
}


# ---------------------------------------------------------------------------
# the door (broker)
# ---------------------------------------------------------------------------

class DoorBroker(threading.Thread):
    """One socket per run. One JSON line in, one JSON line out, one
    request per connection. Validation refuses loud; validated calls
    SPEND before they execute so attempts count."""

    def __init__(self, rid, manifest, sock_path):
        super().__init__(daemon=True)
        self.rid = rid
        self.manifest = manifest
        self.sock_path = sock_path
        self.stop_ev = threading.Event()
        self.refused = 0
        self.refused_capped = False
        self.sock = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
        if os.path.lexists(sock_path):
            os.unlink(sock_path)   # stale node from a dead listener (bug 6)
        self.sock.bind(sock_path)
        self.sock.listen(8)
        self.sock.settimeout(0.5)

    def run(self):
        while not self.stop_ev.is_set():
            try:
                conn, _ = self.sock.accept()
            except socket.timeout:
                continue
            except OSError:
                break
            try:
                self._handle(conn)
            except Exception as e:
                audited("DOOR_INTERNAL", self.rid, "broker error: %s" % e)
            finally:
                try:
                    conn.close()
                except OSError:
                    pass
        try:
            self.sock.close()
        except OSError:
            pass

    def stop(self):
        self.stop_ev.set()
        try:
            self.sock.close()
        except OSError:
            pass
        # close() does not remove a unix socket node; unlink it (bug 6)
        try:
            os.unlink(self.sock_path)
        except FileNotFoundError:
            pass
        except OSError as e:
            audited("DOOR_INTERNAL", self.rid, "door unlink failed: %s" % e)

    # -- protocol -----------------------------------------------------

    def _recv_line(self, conn):
        conn.settimeout(10)
        buf = b""
        while len(buf) <= REQ_LINE_CAP:
            try:
                chunk = conn.recv(4096)
            except socket.timeout:
                raise SkillError("request timed out")
            if not chunk:
                break
            buf += chunk
            if b"\n" in buf:
                line, _, _ = buf.partition(b"\n")
                return line
        raise SkillError("request too large or truncated (cap %d)"
                         % REQ_LINE_CAP)

    def _send(self, conn, obj):
        line = json.dumps(obj, ensure_ascii=True)
        if len(line) > RESP_LINE_CAP:
            line = json.dumps({"ok": False,
                               "error": "response exceeded %d bytes"
                                        % RESP_LINE_CAP})
        try:
            conn.sendall(line.encode("ascii") + b"\n")
        except OSError:
            pass

    def _refuse(self, conn, kind, msg):
        self.refused += 1
        if self.refused > REFUSED_CAP:
            if not self.refused_capped:
                self.refused_capped = True
                audited("DOOR_REFUSED_CAP", self.rid,
                        "refusal cap %d reached; further refusals dropped"
                        % REFUSED_CAP)
            self._send(conn, {"ok": False, "error": "refusal cap reached"})
            return
        audited(kind, self.rid, msg)
        self._send(conn, {"ok": False, "error": msg})

    # -- dispatch -----------------------------------------------------

    def _handle(self, conn):
        try:
            raw = self._recv_line(conn)
        except SkillError as e:
            self._refuse(conn, "DOOR_BADREQ", str(e))
            return
        try:
            req = json.loads(raw.decode("utf-8", "replace"))
            if not isinstance(req, dict):
                raise ValueError("request is not an object")
        except Exception as e:
            self._refuse(conn, "DOOR_BADREQ", "unparseable request: %s" % e)
            return
        skill = req.get("skill", "")
        args = req.get("args", {})
        if not isinstance(args, dict):
            self._refuse(conn, "DOOR_BADREQ", "args must be an object")
            return
        if skill not in KNOWN_SKILLS:
            self._refuse(conn, "DOOR_REFUSED", "unknown skill %r" % skill)
            return
        skills = self.manifest.get("skills", {})
        if skill not in skills:
            self._refuse(conn, "DOOR_REFUSED",
                         "skill %s not granted to %s"
                         % (skill, self.manifest["tool"]))
            return
        cfg = skills[skill]
        # budgets: run cap first, then per-skill cap; both counted from
        # SPENT rows so a crashed broker never undercounts.
        conn_db = db()
        try:
            cur = conn_db.execute(
                "SELECT COUNT(*) FROM door_calls WHERE rid=?", (self.rid,))
            spent_total = cur.fetchone()[0]
            cur = conn_db.execute(
                "SELECT COUNT(*) FROM door_calls WHERE rid=? AND skill=?",
                (self.rid, skill))
            spent_skill = cur.fetchone()[0]
        finally:
            conn_db.close()
        if spent_total >= self.manifest["door_call_cap"]:
            self._refuse(conn, "DOOR_REFUSED",
                         "run door-call cap %d reached"
                         % self.manifest["door_call_cap"])
            return
        max_calls = int(cfg.get("max_calls", 0))
        if spent_skill >= max_calls:
            self._refuse(conn, "DOOR_REFUSED",
                         "%s max_calls %d reached" % (skill, max_calls))
            return
        seq = spent_total + 1
        dgst = digest_obj(args)

        def spend(c):
            c.execute(
                "INSERT INTO door_calls(rid,seq,skill,args_digest,state,"
                "started) VALUES(?,?,?,?,?,?)",
                (self.rid, seq, skill, dgst, "SPENT", now()))

        audited("DOOR_SPEND", self.rid,
                "seq=%d skill=%s digest=%s" % (seq, skill, dgst), fn=spend)
        try:
            result, meta = SKILL_FNS[skill](cfg, args)
            ok, err = 1, None
        except SkillError as e:
            result, meta = None, {}
            ok, err = 0, str(e)
        except Exception as e:
            result, meta = None, {}
            ok, err = 0, "skill internal error: %s" % e

        def close_call(c):
            c.execute(
                "UPDATE door_calls SET state=?, ended=?, ok=?, error=?, "
                "tier=?, model=? WHERE rid=? AND seq=?",
                ("DONE" if ok else "ERROR", now(), ok, err,
                 meta.get("tier"), meta.get("model"), self.rid, seq))

        audited("DOOR_OK" if ok else "DOOR_ERR", self.rid,
                "seq=%d skill=%s %s" % (seq, skill, err or "ok"),
                fn=close_call)
        if ok:
            self._send(conn, {"ok": True, "result": result})
        else:
            self._send(conn, {"ok": False, "error": err})


# ---------------------------------------------------------------------------
# the room
# ---------------------------------------------------------------------------

# CONTAINMENT T2 rev 3 (Tex item 2(c)): the room's launch configuration is
# root-owned. podman is called by absolute path (no PATH lookup), and its
# environment is fixed: CONTAINERS_CONF names the root-owned distro default
# ("all system and user config files are ignored", containers.conf(5)), and
# XDG_CONFIG_HOME points at an empty root-owned dir, so no treebeard-writable
# containers.conf / mounts.conf / registries.conf can change the room.
# Today no such user file exists, so the effective config is unchanged.
# Podman's runtime STORAGE stays treebeard's (recorded containment debt).
ROOM_PODMAN = "/usr/bin/podman"
ROOM_CONTAINERS_CONF = "/usr/share/containers/containers.conf"
ROOM_XDG_CONFIG = "/opt/treebeard/shared/toolbox/podman-config"


def _room_env():
    env = {k: v for k, v in os.environ.items()
           if not k.startswith(("CONTAINERS_", "REGISTRIES_"))}
    env["CONTAINERS_CONF"] = ROOM_CONTAINERS_CONF
    env["XDG_CONFIG_HOME"] = ROOM_XDG_CONFIG
    return env


def _room_cmd(name, tooldir, scratch, doordir, image, tool_args):
    """THE seam. Everything the wall is, in one argv. Callers do not
    change when the flags do."""
    return [
        ROOM_PODMAN, "run", "--rm", "--name", name,
        "--network=none",
        "--userns=keep-id",
        "--security-opt", "no-new-privileges",
        "--cap-drop", "ALL",
        "--pids-limit", ROOM_PIDS,
        "--memory", ROOM_MEM,
        "--cpus", ROOM_CPUS,
        "--read-only", "--tmpfs", ROOM_TMPFS,
        "-v", "%s:/tool:ro" % tooldir,
        "-v", "%s:/scratch:rw" % scratch,
        "-v", "%s:/door:rw" % doordir,
        "-w", "/scratch",
        "-e", "TOOLBOX_DOOR=/door/door.sock",
        image, "python3", "/tool/main.py",
    ] + list(tool_args)


def _new_rid():
    return "r" + hashlib.sha256(
        ("%s-%s-%s" % (now(), os.getpid(), time.monotonic_ns())
         ).encode()).hexdigest()[:11]


def _sweep_dead_doors():
    """Remove door.sock nodes whose listener died without stop() (bug 6:
    toolbox killed mid-run). Predicate is the ledger: only runs with a
    TERMINAL row are touched; PENDING/RUNNING and unknown rids are left
    alone. Only real socket nodes are unlinked. Returns count removed."""
    import stat as _stat
    runs_dir = _p("runs")
    if not os.path.isdir(runs_dir):
        return 0
    conn = db()
    try:
        terminal = {r[0] for r in conn.execute(
            "SELECT rid FROM runs WHERE state NOT IN ('PENDING','RUNNING')")}
    finally:
        conn.close()
    removed = 0
    for rid in os.listdir(runs_dir):
        if rid not in terminal:
            continue
        sp = os.path.join(runs_dir, rid, "door", "door.sock")
        try:
            if not _stat.S_ISSOCK(os.lstat(sp).st_mode):
                continue
            os.unlink(sp)
            removed += 1
        except FileNotFoundError:
            continue
        except OSError:
            continue
    if removed:
        audited("DOOR_SWEEP", None, "removed %d dead door socket(s)" % removed)
    return removed


def run_tool(tool, tool_args):
    check_tool_name(tool)
    m = load_manifest(tool)
    if m is None:
        raise FileNotFoundError(
            "no manifest for %s -- a tool with no manifest cannot run "
            "(enroll it, then allow skills deliberately)" % tool)
    invoker = getpass.getuser()
    if invoker not in m["invokers"]:
        audited("RUN_REFUSED", None,
                "tool=%s invoker=%s not in invokers" % (tool, invoker))
        raise PermissionError(
            "%s may not invoke %s (invokers: %s)"
            % (invoker, tool, ",".join(m["invokers"])))
    tooldir = os.path.join(TOOLKIT(), tool)
    entry = os.path.join(tooldir, "main.py")
    if not os.path.isfile(entry):
        raise FileNotFoundError("tool entry missing: %s" % entry)

    unsafe = os.environ.get("TOOLBOX_UNSAFE_ROOM") == "1"
    room = "UNSAFE-SUBPROCESS" if unsafe else "podman"

    rid = _new_rid()
    _sweep_dead_doors()
    rundir = _p("runs", rid)
    scratch = os.path.join(rundir, "scratch")
    doordir = os.path.join(rundir, "door")
    logsdir = os.path.join(rundir, "logs")
    for d in (scratch, doordir, logsdir,
              os.path.join(scratch, "result")):
        os.makedirs(d, exist_ok=True)

    def create(c):
        c.execute(
            "INSERT INTO runs(rid,tool,argv,state,room,image,created,"
            "invoker) VALUES(?,?,?,?,?,?,?,?)",
            (rid, tool, json.dumps(list(tool_args)), "PENDING", room,
             IMAGE() if not unsafe else "-", now(), invoker))

    audited("RUN_CREATE", rid, "tool=%s room=%s" % (tool, room), fn=create)
    write_card()

    sock_path = os.path.join(doordir, "door.sock")
    broker = DoorBroker(rid, m, sock_path)
    broker.start()

    name = "tb-%s" % rid
    if unsafe:
        cmd = [sys.executable, entry] + list(tool_args)
        env = dict(os.environ)
        env["TOOLBOX_DOOR"] = sock_path
        popen_kw = {"cwd": scratch, "env": env}
    else:
        cmd = _room_cmd(name, tooldir, scratch, doordir, IMAGE(), tool_args)
        popen_kw = {"env": _room_env()}

    out_f = open(os.path.join(logsdir, "stdout.txt"), "wb")
    err_f = open(os.path.join(logsdir, "stderr.txt"), "wb")
    started = now()

    def start(c):
        c.execute("UPDATE runs SET state='RUNNING', started=? WHERE rid=?",
                  (started, rid))

    audited("RUN_START", rid, "cmd[0..2]=%s" % " ".join(cmd[:3]), fn=start)
    write_card()

    timed_out = False
    try:
        proc = subprocess.Popen(cmd, stdout=out_f, stderr=err_f, **popen_kw)
    except FileNotFoundError as e:
        # podman missing (or python in unsafe mode). Fail-closed and loud.
        broker.stop()
        out_f.close()
        err_f.close()

        def failc(c):
            c.execute(
                "UPDATE runs SET state='FAILED', ended=?, note=? WHERE rid=?",
                (now(), "launch failed: %s" % e, rid))

        audited("RUN_FAILED", rid, "launch failed: %s" % e, fn=failc)
        write_card()
        raise RuntimeError("room launch failed (%s). If this is podman "
                           "missing, the wall does not exist yet -- "
                           "install it; the toolbox will not fall back "
                           "silently." % e)

    wall = int(m["wall_seconds"])
    try:
        proc.wait(timeout=wall)
    except subprocess.TimeoutExpired:
        timed_out = True
        if unsafe:
            proc.terminate()
        else:
            subprocess.run([ROOM_PODMAN, "kill", name], env=_room_env(),
                           stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL)
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait()
    finally:
        out_f.close()
        err_f.close()
        broker.stop()
        broker.join(timeout=5)

    # collect result files (names + hashes) for the run record
    results = []
    resdir = os.path.join(scratch, "result")
    if os.path.isdir(resdir):
        for base, _, files in os.walk(resdir):
            for fn in sorted(files):
                path = os.path.join(base, fn)
                rel = os.path.relpath(path, resdir)
                try:
                    results.append({"file": rel,
                                    "sha256": sha256_file(path),
                                    "bytes": os.path.getsize(path)})
                except OSError:
                    results.append({"file": rel, "sha256": "unreadable",
                                    "bytes": -1})

    if timed_out:
        state, note = "TIMEOUT", "killed at wall %ds" % wall
    elif proc.returncode == 0:
        state, note = "DONE", ""
    else:
        state, note = "FAILED", "exit %s" % proc.returncode

    def close(c):
        c.execute(
            "UPDATE runs SET state=?, ended=?, exit_code=?, note=? "
            "WHERE rid=?",
            (state, now(), proc.returncode, json.dumps(
                {"note": note, "results": results}), rid))
        # any call SPENT but never closed gets closed as ABANDONED in the
        # same transaction as the run close -- the ledger never dangles.
        c.execute(
            "UPDATE door_calls SET state='ABANDONED', ended=?, ok=0, "
            "error='abandoned at run close' WHERE rid=? AND state='SPENT'",
            (now(), rid))

    audited("RUN_%s" % state, rid, note or "clean exit", fn=close)
    write_card()
    return rid, state, results


def cmd_run(tool, tool_args):
    rid, state, results = run_tool(tool, tool_args)
    print("run %s: %s -> %s" % (rid, tool, state))
    for r in results:
        print("  result %s (%d bytes) sha256=%s"
              % (r["file"], r["bytes"], r["sha256"]))
    print("logs: %s" % _p("runs", rid, "logs"))
    if state != "DONE":
        sys.exit(1)


# ---------------------------------------------------------------------------
# probe (real-wall proof; the anvil probe_tool equivalent)
# ---------------------------------------------------------------------------

PROBE_MAIN = r'''# probe payload -- written by the toolbox, runs INSIDE the room.
# Proves the wall holds and the door works. Prints one JSON line.
import json, os, socket

checks = []
def add(name, ok, info=""):
    checks.append({"name": name, "ok": bool(ok), "info": str(info)[:200]})

# 1. a host-only path must be invisible (the wall)
hostpath = os.environ.get("PROBE_HOSTPATH", "/opt/treebeard/shared")
try:
    open(hostpath if os.path.isfile(hostpath)
         else os.path.join(hostpath, "x"), "rb").close()
    add("host path invisible", False, "READ SUCCEEDED: wall breached")
except OSError as e:
    add("host path invisible", True, type(e).__name__)

# 2. network must be unreachable (--network=none)
try:
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    s.settimeout(3)
    s.connect(("1.1.1.1", 53))
    s.close()
    add("network dark", False, "CONNECT SUCCEEDED: egress exists")
except OSError as e:
    add("network dark", True, type(e).__name__)

# 3. the tool mount must be read-only
try:
    open("/tool/write_probe", "w").close()
    add("tool mount read-only", False, "WRITE SUCCEEDED")
except OSError as e:
    add("tool mount read-only", True, type(e).__name__)

# 4. the door must answer (state: toolbox card)
try:
    c = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
    c.settimeout(10)
    c.connect(os.environ["TOOLBOX_DOOR"])
    c.sendall((json.dumps({"v": 1, "skill": "state",
                           "args": {"card": "toolbox"}}) + "\n").encode())
    buf = b""
    while b"\n" not in buf:
        ch = c.recv(4096)
        if not ch:
            break
        buf += ch
    c.close()
    resp = json.loads(buf.decode("utf-8", "replace"))
    add("door answers", resp.get("ok") is True,
        resp.get("error", "ok"))
except Exception as e:
    add("door answers", False, "%s: %s" % (type(e).__name__, e))

# 5. scratch writes work, temp-then-rename, plus a cleanup target
os.makedirs("/scratch/result", exist_ok=True)
os.makedirs("/scratch/cleanup_check", exist_ok=True)
try:
    with open("/scratch/result/probe.json.tmp", "w") as f:
        f.write(json.dumps({"checks": checks}))
    os.replace("/scratch/result/probe.json.tmp", "/scratch/result/probe.json")
    with open("/scratch/cleanup_check/f.txt", "w") as f:
        f.write("owned-by-keep-id\n")
    add("scratch writable", True, "")
except OSError as e:
    add("scratch writable", False, str(e))

print(json.dumps({"checks": checks}))
raise SystemExit(0 if all(c["ok"] for c in checks) else 1)
'''


def cmd_probe():
    if os.environ.get("TOOLBOX_UNSAFE_ROOM") == "1":
        raise RuntimeError("probe proves the REAL wall; unset "
                           "TOOLBOX_UNSAFE_ROOM")
    # ephemeral tool dir under the data tree (not the toolkit; the probe
    # is the toolbox proving itself, not a promoted tool)
    pdir = _p("staging", "probe-tool")
    if os.path.isdir(pdir):
        shutil.rmtree(pdir)
    os.makedirs(pdir)
    with open(os.path.join(pdir, "main.py"), "w", encoding="ascii") as f:
        f.write(PROBE_MAIN)
    manifest = {
        "tool": "__probe__",
        "invokers": [getpass.getuser()],
        "wall_seconds": 60,
        "door_call_cap": 4,
        "skills": {"state": {"cards": ["toolbox"], "max_calls": 2}},
    }
    write_card()  # make sure the toolbox card exists for check 4
    rid = _new_rid()
    rundir = _p("runs", rid)
    scratch = os.path.join(rundir, "scratch")
    doordir = os.path.join(rundir, "door")
    logsdir = os.path.join(rundir, "logs")
    for d in (scratch, doordir, logsdir):
        os.makedirs(d, exist_ok=True)

    def create(c):
        c.execute(
            "INSERT INTO runs(rid,tool,argv,state,room,image,created,"
            "invoker) VALUES(?,?,?,?,?,?,?,?)",
            (rid, "__probe__", "[]", "RUNNING", "podman", IMAGE(), now(),
             getpass.getuser()))

    audited("PROBE_START", rid, "image=%s" % IMAGE(), fn=create)
    sock_path = os.path.join(doordir, "door.sock")
    broker = DoorBroker(rid, manifest, sock_path)
    broker.start()
    name = "tb-%s" % rid
    cmd = _room_cmd(name, pdir, scratch, doordir, IMAGE(), [])
    # pass the host-only path in via -e (insert before image arg)
    idx = cmd.index(IMAGE())
    cmd[idx:idx] = ["-e", "PROBE_HOSTPATH=%s" % _p("ledger.db")]
    out_f = open(os.path.join(logsdir, "stdout.txt"), "wb")
    err_f = open(os.path.join(logsdir, "stderr.txt"), "wb")
    try:
        proc = subprocess.Popen(cmd, stdout=out_f, stderr=err_f, env=_room_env())
        try:
            proc.wait(timeout=90)
        except subprocess.TimeoutExpired:
            subprocess.run([ROOM_PODMAN, "kill", name], env=_room_env(),
                           stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL)
            proc.wait()
    finally:
        out_f.close()
        err_f.close()
        broker.stop()
        broker.join(timeout=5)

    stdout = open(os.path.join(logsdir, "stdout.txt"), "rb").read()
    checks = []
    for line in stdout.decode("utf-8", "replace").splitlines():
        line = line.strip()
        if line.startswith("{"):
            try:
                checks = json.loads(line).get("checks", [])
            except ValueError:
                pass
    # host-side checks: keep-id ownership + rmtree of container-written dir
    cleanup = os.path.join(scratch, "cleanup_check")
    own_ok = False
    rm_ok = False
    if os.path.isdir(cleanup):
        try:
            st = os.stat(os.path.join(cleanup, "f.txt"))
            own_ok = (st.st_uid == os.getuid())
        except OSError:
            own_ok = False
        try:
            shutil.rmtree(cleanup)
            rm_ok = True
        except OSError:
            rm_ok = False
    checks.append({"name": "keep-id ownership (host)", "ok": own_ok,
                   "info": ""})
    checks.append({"name": "host rmtree of room writes", "ok": rm_ok,
                   "info": ""})
    passed = sum(1 for c in checks if c.get("ok"))
    total = len(checks)

    def close(c):
        c.execute(
            "UPDATE runs SET state=?, ended=?, exit_code=? WHERE rid=?",
            ("DONE" if passed == total and total >= 7 else "FAILED",
             now(), proc.returncode, rid))

    audited("PROBE_DONE", rid, "%d/%d" % (passed, total), fn=close)
    write_card()
    for c in checks:
        print("%-32s %s %s" % (c.get("name", "?"),
                               "ok" if c.get("ok") else "FAIL",
                               c.get("info", "")))
    if err := open(os.path.join(logsdir, "stderr.txt"), "rb").read():
        if passed != total:
            print("--- stderr tail ---")
            print(err.decode("utf-8", "replace")[-1500:])
    print("PROBE %d/%d %s" % (passed, total,
                              "PASS" if passed == total and total >= 7
                              else "FAIL"))
    if passed != total or total < 7:
        sys.exit(1)


# ---------------------------------------------------------------------------
# knowledge skills (docs; not door services)
# ---------------------------------------------------------------------------

TOOL_CONTRACT_MD = """# TOOL CONTRACT (toolbox room)

Every promoted tool the toolbox runs must obey this. It is the room's
side of the bargain; the manifest is treebeard's.

ENTRY   The room executes `python3 /tool/main.py <args...>`. Your tool
        directory is mounted read-only at /tool. Working directory is
        /scratch (writable). /tmp is a small tmpfs.

DOOR    The only way out is the socket at $TOOLBOX_DOOR. Protocol: open
        a connection, send ONE JSON line, read ONE JSON line, close.
          {"v":1, "skill":"state", "args":{"card":"meeseeks"}}
          {"v":1, "skill":"router", "args":{"system":"...", "user":"...",
            "tier":"free", "max_tokens":800}}
          {"v":1, "skill":"web_search", "args":{"query":"...", "count":3}}
        Responses: {"ok":true, "result":...} or {"ok":false, "error":"..."}.
        A skill your manifest does not grant is refused; budgets are
        enforced host-side and attempts count. Handle "ok": false --
        refusal is a normal answer, not a crash.

RESULT  Write durable output under /scratch/result/ using temp-then-
        rename (write name.tmp, os.replace to name). The host records
        every result file with its sha256. stdout/stderr are captured to
        the run's logs; print progress there.

LIMITS  No network. No filesystem outside /tool (ro), /scratch, /tmp.
        A wall clock kills the run at manifest wall_seconds; if your
        work can die mid-write, temp-then-rename is what keeps the
        result pool uncorrupted.

STYLE   Plain ASCII. Comment intent, not noise. Standard library only
        unless your task said otherwise (the room image carries the
        document libraries below when it does).
"""

PDF_SKILL_MD = """# SKILL: pdf (knowledge skill -- library + doc, not a door service)

Library: pypdf (read/merge/split/extract), reportlab (create).
Install (host venv or room image): pip install pypdf reportlab

READ / EXTRACT
    from pypdf import PdfReader
    r = PdfReader("in.pdf")
    text = "\\n".join(page.extract_text() or "" for page in r.pages)

MERGE / SPLIT
    from pypdf import PdfWriter, PdfReader
    w = PdfWriter()
    for src in ("a.pdf", "b.pdf"):
        for page in PdfReader(src).pages:
            w.add_page(page)
    with open("out.pdf.tmp", "wb") as f: w.write(f)
    import os; os.replace("out.pdf.tmp", "out.pdf")

CREATE
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    c = canvas.Canvas("out.pdf.tmp", pagesize=letter)
    c.setFont("Helvetica", 11)
    y = 750
    for line in lines:
        c.drawString(72, y, line); y -= 14
        if y < 72: c.showPage(); c.setFont("Helvetica", 11); y = 750
    c.save()
    import os; os.replace("out.pdf.tmp", "out.pdf")

DISCIPLINES
- temp-then-rename on every output (the wall clock can kill mid-write).
- extract_text() returns None on image-only pages; guard it. There is no
  OCR in this stack; an image-only PDF is a CANT, not a retry.
- Ship a test_*.py: assert page count, assert a known string extracts.
"""

XLSX_SKILL_MD = """# SKILL: xlsx (knowledge skill -- library + doc, not a door service)

Library: openpyxl.
Install (host venv or room image): pip install openpyxl

READ
    from openpyxl import load_workbook
    wb = load_workbook("in.xlsx", data_only=True)  # data_only: values,
    ws = wb.active                                 # not formula strings
    rows = [[c.value for c in row] for row in ws.iter_rows()]

WRITE
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "data"
    ws.append(["name", "count"])          # header row
    for name, count in items: ws.append([name, count])
    wb.save("out.xlsx.tmp")
    import os; os.replace("out.xlsx.tmp", "out.xlsx")

FORMULAS / FORMAT
    ws["C2"] = "=A2*B2"                   # written as text; Excel computes
    ws.column_dimensions["A"].width = 24
    from openpyxl.styles import Font
    for c in ws[1]: c.font = Font(bold=True)

DISCIPLINES
- data_only=True reads cached values; a file never opened by Excel has
  none. Compute in Python when the number must be real.
- temp-then-rename on save; .xlsx is a zip and a torn zip is garbage.
- Ship a test_*.py: reload the written file, assert cell values.
"""

PPTX_SKILL_MD = """# SKILL: pptx (knowledge skill -- library + doc, not a door service)

Library: python-pptx.
Install (host venv or room image): pip install python-pptx

CREATE
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title+content
    slide.shapes.title.text = "Title"
    body = slide.placeholders[1].text_frame
    body.text = "first line"
    p = body.add_paragraph(); p.text = "second line"; p.level = 1
    prs.save("out.pptx.tmp")
    import os; os.replace("out.pptx.tmp", "out.pptx")

READ / EDIT
    prs = Presentation("in.pptx")
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)

DISCIPLINES
- Layout indexes vary by template; discover placeholders by iterating
  slide.placeholders and reading .placeholder_format.idx, do not guess.
- temp-then-rename on save; .pptx is a zip, same torn-zip failure as xlsx.
- Ship a test_*.py: reopen the deck, assert slide count and title text.
"""

SKILL_DOCS = {
    "tool_contract": TOOL_CONTRACT_MD,
    "pdf": PDF_SKILL_MD,
    "xlsx": XLSX_SKILL_MD,
    "pptx": PPTX_SKILL_MD,
}


# ---------------------------------------------------------------------------
# commands: init / status / listings
# ---------------------------------------------------------------------------

def cmd_init():
    for d in ("manifests", "config", "skills", "runs", "status", "staging"):
        os.makedirs(_p(d), exist_ok=True)
    conn = db()
    try:
        with conn:
            conn.executescript(SCHEMA)
            event(conn, "INIT", None, "toolbox initialized at %s" % ROOT())
    finally:
        conn.close()
    if not os.path.exists(cards_path()):
        seed_cards()
    for name, text in SKILL_DOCS.items():
        require_ascii(text, "skill doc %s" % name)
        path = _p("skills", "%s.md" % name)
        if not os.path.exists(path):
            atomic_write(path, text)
    write_card()
    print("toolbox initialized at %s" % ROOT())
    print("cards registry: %s (treebeard-owned; edit paths as truth requires)"
          % cards_path())
    print("skill docs: %s" % ", ".join(sorted(SKILL_DOCS)))
    print("status card: %s" % _p("status", "toolbox.txt"))
    print("NOTE: ownership pass at deploy -- treebeard owns the tree;")
    print("NOTE: the wall requires podman + image %s; gate the walk on"
          % IMAGE())
    print("      `podman --version`, then `toolbox probe` before any run.")
    print("NOTE: confirm router result attribute names (text/tier/model)")
    print("      against the deployed router.py before the first live")
    print("      router-skill call.")


def cmd_status(tool=None):
    card = _p("status", "toolbox.txt")
    if os.path.exists(card):
        sys.stdout.write(open(card, "r", encoding="ascii",
                              errors="replace").read())
    conn = db()
    try:
        if tool:
            rows = conn.execute(
                "SELECT rid,state,room,created,ended,exit_code FROM runs "
                "WHERE tool=? ORDER BY created DESC LIMIT 10",
                (tool,)).fetchall()
            print("--- %s runs ---" % tool)
            for r in rows:
                print("  %s [%s] room=%s %s -> %s exit=%s"
                      % (r[0], r[1], r[2], r[3], r[4] or "-", r[5]))
            if rows:
                rid = rows[0][0]
                calls = conn.execute(
                    "SELECT seq,skill,state,ok,error,tier,model FROM "
                    "door_calls WHERE rid=? ORDER BY seq", (rid,)).fetchall()
                print("--- latest run %s door calls ---" % rid)
                for c in calls:
                    print("  #%d %s [%s] ok=%s tier=%s model=%s %s"
                          % (c[0], c[1], c[2], c[3], c[5] or "-",
                             c[6] or "-", c[4] or ""))
                for stream in ("stdout", "stderr"):
                    lp = _p("runs", rid, "logs", "%s.txt" % stream)
                    if os.path.exists(lp):
                        data = open(lp, "rb").read()[-LOG_SHOW_CAP:]
                        if data.strip():
                            print("--- %s tail ---" % stream)
                            print(data.decode("utf-8", "replace"))
    finally:
        conn.close()


def cmd_skills():
    sk = _p("skills")
    print("knowledge skill docs at %s" % sk)
    print("(callers prepend the relevant doc to prompts; these are not")
    print(" door services -- the library plus the doc IS the skill)")
    if os.path.isdir(sk):
        for n in sorted(os.listdir(sk)):
            if n.endswith(".md"):
                print("  %s" % n)


def cmd_cards():
    reg = load_cards()
    print("state-skill card registry (%s):" % cards_path())
    for name in sorted(reg):
        path = reg[name]
        mark = "ok" if os.path.exists(path) else "MISSING"
        print("  %-10s %-7s %s" % (name, mark, path))


# ---------------------------------------------------------------------------
# selftest
# ---------------------------------------------------------------------------

FAKE_ROUTER = '''# fake router for the toolbox selftest -- scripted, no network
import os
class _G(object):
    pass
def _tier_names():
    return ["local", "free", "cheap", "frontier"]
def route(system, user, force_tier=None, **kw):
    g = _G()
    if os.environ.get("FAKE_ROUTER_FAIL") == "1":
        g.ok = False
        g.error = "scripted failure"
        g.text = ""
        g.tier = force_tier or "free"
        g.model = None
        return g
    if force_tier not in _tier_names():
        raise ValueError("force_tier %r is not a known tier" % force_tier)
    g.ok = True
    g.error = None
    g.text = "FAKE[%s] %s" % (force_tier, user[:40])
    g.tier = force_tier
    g.model = "fake-model-1"
    return g
'''

TOOL_OK = r'''# selftest tool: one state call, one router call, writes a result
import json, os, socket, sys

def door(req):
    c = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
    c.settimeout(15)
    c.connect(os.environ["TOOLBOX_DOOR"])
    c.sendall((json.dumps(req) + "\n").encode())
    buf = b""
    while b"\n" not in buf:
        ch = c.recv(4096)
        if not ch:
            break
        buf += ch
    c.close()
    return json.loads(buf.decode("utf-8", "replace"))

r1 = door({"v": 1, "skill": "state", "args": {"card": "toolbox"}})
print("STATE_OK=%s" % r1.get("ok"))
r2 = door({"v": 1, "skill": "router",
           "args": {"system": "s", "user": "hello world", "tier": "free"}})
print("ROUTER_OK=%s TEXT=%s MODEL=%s"
      % (r2.get("ok"), r2.get("result", {}).get("text", ""),
         r2.get("result", {}).get("model", "")))
r3 = door({"v": 1, "skill": "web_search", "args": {"query": "q", "count": 2}})
print("SEARCH_OK=%s N=%d" % (r3.get("ok"),
                             len(r3.get("result", {}).get("results", []))))
os.makedirs("result", exist_ok=True)
with open("result/out.txt.tmp", "w") as f:
    f.write("ok\n")
os.replace("result/out.txt.tmp", "result/out.txt")
sys.exit(0)
'''

TOOL_DENY = r'''# selftest tool: asks for a skill it was never granted
import json, os, socket, sys
c = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
c.settimeout(15)
c.connect(os.environ["TOOLBOX_DOOR"])
c.sendall((json.dumps({"v": 1, "skill": "router",
                       "args": {"user": "x", "tier": "free"}}) + "\n"
           ).encode())
buf = b""
while b"\n" not in buf:
    ch = c.recv(4096)
    if not ch:
        break
    buf += ch
c.close()
resp = json.loads(buf.decode("utf-8", "replace"))
print("DENIED=%s ERR=%s" % (resp.get("ok") is False,
                            resp.get("error", "")[:60]))
sys.exit(0)
'''

TOOL_HOG = r'''# selftest tool: hammers a skill until budget blocks, then bad calls
import json, os, socket, sys

def door(payload_line):
    c = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
    c.settimeout(15)
    c.connect(os.environ["TOOLBOX_DOOR"])
    c.sendall(payload_line)
    buf = b""
    while b"\n" not in buf:
        ch = c.recv(4096)
        if not ch:
            break
        buf += ch
    c.close()
    return json.loads(buf.decode("utf-8", "replace"))

okc = 0
refc = 0
for i in range(5):
    r = door((json.dumps({"v": 1, "skill": "state",
                          "args": {"card": "toolbox"}}) + "\n").encode())
    if r.get("ok"):
        okc += 1
    else:
        refc += 1
g = door(b"this is not json\n")
u = door((json.dumps({"v": 1, "skill": "shell",
                      "args": {"cmd": "rm -rf /"}}) + "\n").encode())
print("OK=%d REF=%d GARBAGE_REFUSED=%s UNKNOWN_REFUSED=%s"
      % (okc, refc, g.get("ok") is False, u.get("ok") is False))
sys.exit(0)
'''

TOOL_SLEEP = r'''# selftest tool: outlives the wall clock on purpose
import time
time.sleep(60)
'''

TOOL_FAIL = r'''# selftest tool: exits nonzero
import sys
sys.exit(3)
'''

TOOL_RFAIL = r'''# selftest tool: one router call that the fake router fails
import json, os, socket, sys
c = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
c.settimeout(15)
c.connect(os.environ["TOOLBOX_DOOR"])
c.sendall((json.dumps({"v": 1, "skill": "router",
                       "args": {"user": "x", "tier": "free"}}) + "\n"
           ).encode())
buf = b""
while b"\n" not in buf:
    ch = c.recv(4096)
    if not ch:
        break
    buf += ch
c.close()
resp = json.loads(buf.decode("utf-8", "replace"))
print("RFAIL_ERRORED=%s" % (resp.get("ok") is False))
sys.exit(0)
'''


def cmd_selftest():
    tmp = tempfile.mkdtemp(prefix="toolbox-selftest-")
    root = os.path.join(tmp, "data")
    kit = os.path.join(tmp, "toolkit")
    lib = os.path.join(tmp, "lib")
    os.makedirs(kit)
    os.makedirs(lib)
    with open(os.path.join(lib, "router.py"), "w", encoding="ascii") as f:
        f.write(FAKE_ROUTER)
    fake_search = os.path.join(tmp, "search.json")
    with open(fake_search, "w", encoding="ascii") as f:
        json.dump({"results": [
            {"title": "T1", "url": "https://e/1", "description": "D1"},
            {"title": "T2", "url": "https://e/2", "description": "D2"},
            {"title": "T3", "url": "https://e/3", "description": "D3"},
        ]}, f)
    me = getpass.getuser()
    os.environ["TOOLBOX_ROOT"] = root
    os.environ["TOOLBOX_LIB"] = lib
    os.environ["TOOLKIT_ROOT"] = kit
    os.environ["TOOLBOX_GATE_USER"] = me
    os.environ["TOOLBOX_UNSAFE_ROOM"] = "1"
    os.environ["TOOLBOX_SEARCH_FAKE"] = fake_search
    os.environ.pop("FAKE_ROUTER_FAIL", None)

    passed = [0]
    failed = [0]

    def check(name, cond):
        print("%-46s %s" % (name, "ok" if cond else "FAIL"))
        (passed if cond else failed)[0] += 1

    def mktool(name, source):
        d = os.path.join(kit, name)
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, "main.py"), "w", encoding="ascii") as f:
            f.write(source)

    def events_of(kind, rid=None):
        conn = db()
        try:
            if rid:
                return conn.execute(
                    "SELECT detail FROM events WHERE kind=? AND rid=?",
                    (kind, rid)).fetchall()
            return conn.execute(
                "SELECT detail FROM events WHERE kind=?", (kind,)).fetchall()
        finally:
            conn.close()

    def run_row(rid):
        conn = db()
        try:
            return conn.execute(
                "SELECT state,exit_code,room,note FROM runs WHERE rid=?",
                (rid,)).fetchone()
        finally:
            conn.close()

    def calls_of(rid):
        conn = db()
        try:
            return conn.execute(
                "SELECT seq,skill,state,ok,error,tier,model FROM door_calls "
                "WHERE rid=? ORDER BY seq", (rid,)).fetchall()
        finally:
            conn.close()

    # --- init ---------------------------------------------------------
    cmd_init()
    check("init creates db", os.path.exists(_p("ledger.db")))
    check("init writes card", os.path.exists(_p("status", "toolbox.txt")))
    check("init seeds cards.json",
          set(load_cards().keys()) == {"toolbox", "meeseeks", "anvil"})
    check("init writes 4 skill docs",
          all(os.path.exists(_p("skills", "%s.md" % n))
              for n in ("tool_contract", "pdf", "xlsx", "pptx")))
    docs_ascii = True
    for n in SKILL_DOCS:
        try:
            open(_p("skills", "%s.md" % n), encoding="ascii").read()
        except UnicodeDecodeError:
            docs_ascii = False
    check("skill docs are ASCII", docs_ascii)

    # --- names / paths --------------------------------------------------
    bad = False
    try:
        check_tool_name("../evil")
    except ValueError:
        bad = True
    check("bad tool name refused", bad)
    bad = False
    try:
        run_tool("ghost", [])
    except FileNotFoundError:
        bad = True
    check("run without manifest refused", bad)

    # --- manifests ------------------------------------------------------
    mktool("t_ok", TOOL_OK)
    bad = False
    try:
        cmd_enroll("nosuchtool")
    except FileNotFoundError:
        bad = True
    check("enroll refuses missing tool dir", bad)
    cmd_enroll("t_ok")
    m = load_manifest("t_ok")
    check("enroll writes deny-all manifest",
          m is not None and m["skills"] == {} and m["invokers"] == [me])
    os.environ["TOOLBOX_GATE_USER"] = "someone-else"
    bad = False
    try:
        cmd_enroll("t_ok")
    except PermissionError:
        bad = True
    check("gate blocks wrong user", bad)
    os.environ["TOOLBOX_GATE_USER"] = me
    bad = False
    try:
        cmd_allow("t_ok", "shell", ["max_calls=1"])
    except ValueError:
        bad = True
    check("allow refuses unknown skill", bad)
    bad = False
    try:
        cmd_allow("t_ok", "router", ["max_calls=2"])
    except ValueError:
        bad = True
    check("router grant demands explicit tiers", bad)
    cmd_allow("t_ok", "router", ["tiers=free,cheap", "max_calls=2",
                                 "max_tokens=500"])
    cmd_allow("t_ok", "state", ["cards=toolbox", "max_calls=3"])
    cmd_allow("t_ok", "web_search", ["max_calls=2", "max_results=2"])
    m = load_manifest("t_ok")
    check("allow writes granted skills",
          set(m["skills"]) == {"router", "state", "web_search"}
          and m["skills"]["router"]["tiers"] == ["free", "cheap"])
    check("allow audited", len(events_of("ALLOW")) == 3)

    # --- happy run (fake room, real door) ------------------------------
    rid, state, results = run_tool("t_ok", [])
    row = run_row(rid)
    check("t_ok run DONE", state == "DONE" and row[0] == "DONE")
    check("run recorded UNSAFE room loud", row[2] == "UNSAFE-SUBPROCESS")
    check("result file recorded",
          len(results) == 1 and results[0]["file"] == "out.txt"
          and len(results[0]["sha256"]) == 64)
    out = open(_p("runs", rid, "logs", "stdout.txt"), "rb").read().decode()
    check("state skill answered in room", "STATE_OK=True" in out)
    check("router skill answered with fake text",
          "ROUTER_OK=True" in out and "FAKE[free]" in out
          and "MODEL=fake-model-1" in out)
    check("search skill returned capped results",
          "SEARCH_OK=True N=2" in out)
    calls = calls_of(rid)
    check("three door calls spent and closed",
          len(calls) == 3 and all(c[2] == "DONE" and c[3] == 1
                                  for c in calls))
    check("router call carries tier and model",
          any(c[1] == "router" and c[5] == "free" and c[6] == "fake-model-1"
              for c in calls))
    check("spend audited before outcome",
          len(events_of("DOOR_SPEND", rid)) == 3
          and len(events_of("DOOR_OK", rid)) == 3)
    check("run lifecycle audited",
          len(events_of("RUN_CREATE", rid)) == 1
          and len(events_of("RUN_START", rid)) == 1
          and len(events_of("RUN_DONE", rid)) == 1)

    # --- deny-by-default -------------------------------------------------
    mktool("t_deny", TOOL_DENY)
    cmd_enroll("t_deny")
    rid, state, _ = run_tool("t_deny", [])
    out = open(_p("runs", rid, "logs", "stdout.txt"), "rb").read().decode()
    check("ungranted skill refused in-room", "DENIED=True" in out
          and "not granted" in out)
    check("refusal spends nothing", len(calls_of(rid)) == 0)
    check("refusal audited",
          len(events_of("DOOR_REFUSED", rid)) == 1)

    # --- budgets / bad requests -----------------------------------------
    mktool("t_hog", TOOL_HOG)
    cmd_enroll("t_hog")
    cmd_allow("t_hog", "state", ["cards=toolbox", "max_calls=3"])
    rid, state, _ = run_tool("t_hog", [])
    out = open(_p("runs", rid, "logs", "stdout.txt"), "rb").read().decode()
    check("skill max_calls blocks at 3",
          "OK=3 REF=2" in out)
    check("garbage and unknown skill refused",
          "GARBAGE_REFUSED=True" in out and "UNKNOWN_REFUSED=True" in out)
    check("spent rows match allowed calls", len(calls_of(rid)) == 3)
    check("bad request audited",
          len(events_of("DOOR_BADREQ", rid)) == 1)

    # run-level door cap: reuse t_hog with cap 1
    m = load_manifest("t_hog")
    m["door_call_cap"] = 1
    save_manifest(m)
    rid, state, _ = run_tool("t_hog", [])
    out = open(_p("runs", rid, "logs", "stdout.txt"), "rb").read().decode()
    check("run door-call cap blocks fail-closed", "OK=1 REF=4" in out)
    check("cap refusals audited",
          len(events_of("DOOR_REFUSED", rid)) >= 4)

    # --- invokers gate -----------------------------------------------------
    m = load_manifest("t_ok")
    m["invokers"] = ["nobody-real"]
    save_manifest(m)
    bad = False
    try:
        run_tool("t_ok", [])
    except PermissionError:
        bad = True
    check("invokers list blocks the run", bad)
    check("blocked run audited", len(events_of("RUN_REFUSED")) == 1)
    m["invokers"] = [me]
    save_manifest(m)

    # --- router failure path ---------------------------------------------
    mktool("t_rfail", TOOL_RFAIL)
    cmd_enroll("t_rfail")
    cmd_allow("t_rfail", "router", ["tiers=free", "max_calls=1"])
    os.environ["FAKE_ROUTER_FAIL"] = "1"
    rid, state, _ = run_tool("t_rfail", [])
    os.environ.pop("FAKE_ROUTER_FAIL", None)
    out = open(_p("runs", rid, "logs", "stdout.txt"), "rb").read().decode()
    calls = calls_of(rid)
    check("router failure reaches the tool cleanly",
          "RFAIL_ERRORED=True" in out)
    check("failed call still SPENT and closed ERROR",
          len(calls) == 1 and calls[0][2] == "ERROR" and calls[0][3] == 0
          and "scripted failure" in (calls[0][4] or ""))
    check("failure audited", len(events_of("DOOR_ERR", rid)) == 1)

    # tier not granted at skill layer
    err = None
    try:
        skill_router({"tiers": ["free"], "max_tokens": 100},
                     {"user": "x", "tier": "frontier"})
    except SkillError as e:
        err = str(e)
    check("router refuses ungranted tier", err is not None
          and "not granted" in err)

    # --- state skill edges --------------------------------------------------
    err = None
    try:
        skill_state({"cards": ["toolbox"]}, {"card": "meeseeks"})
    except SkillError as e:
        err = str(e)
    check("state refuses ungranted card", err is not None)
    # register a card whose path cannot exist under the temp root, so this
    # tests the skill's read-failure path and not whether the box happens
    # to have a real Anvil deployed (it does; the old absolute path made
    # this assertion depend on box state instead of code).
    ghost_path = os.path.join(tmp, "no-such-dir", "ghost.txt")
    reg = load_cards()
    reg["ghost"] = ghost_path
    atomic_write(cards_path(),
                 json.dumps(reg, indent=2, sort_keys=True) + "\n")
    err = None
    try:
        skill_state({"cards": ["ghost"]}, {"card": "ghost"})
    except SkillError as e:
        err = str(e)
    check("state reports unreadable card cleanly",
          err is not None and "cannot read" in err)
    res, _ = skill_state({"cards": ["toolbox"]}, {"card": "toolbox"})
    check("state returns registered card text",
          res["card"] == "toolbox" and "TOOLBOX --" in res["text"])

    # --- search skill edges ---------------------------------------------------
    # cap resolves min(count=99, manifest=5, ceiling=8) -> 5; the fake
    # file carries 3, so 3 come back. Then a manifest of 2 must cut to 2.
    res, _ = skill_web_search({"max_results": 5}, {"query": "q", "count": 99})
    check("search cap above supply returns all 3",
          len(res["results"]) == 3)
    res, _ = skill_web_search({"max_results": 2}, {"query": "q", "count": 99})
    check("search manifest cap cuts results to 2",
          len(res["results"]) == 2)
    saved = os.environ.pop("TOOLBOX_SEARCH_FAKE")
    # 2026-09-01 (PO ruling, test-execution story): this check was
    # environment-sensitive -- it asserted the dark error whatever the
    # outer env held, so any runtime carrying a provider key (Treebeard's
    # ask.py runtime does) turned it permanently red. Precondition-aware:
    # the dark invariant is asserted under a CONTROLLED keyless env, so
    # it holds everywhere; and when the outer env really has a key, the
    # keyed path must NOT report dark -- success or a non-dark failure
    # (auth, network) both prove the key was picked up. The keyed branch
    # makes one live search attempt.
    _saved_keys = {k: os.environ.pop(k) for k in SEARCH_KEY_ENVS
                   if os.environ.get(k)}
    err = None
    try:
        skill_web_search({"max_results": 3}, {"query": "q"})
    except SkillError as e:
        err = str(e)
    os.environ.update(_saved_keys)
    check("search ships dark without a key",
          err is not None and "dark" in err)
    if _saved_keys:
        err2 = None
        try:
            skill_web_search({"max_results": 3}, {"query": "q"})
        except SkillError as e:
            err2 = str(e)
        check("search with a key does not report dark",
              err2 is None or "dark" not in err2)
    os.environ["TOOLBOX_SEARCH_FAKE"] = saved

    # --- timeout and failure runs ---------------------------------------------
    mktool("t_sleep", TOOL_SLEEP)
    cmd_enroll("t_sleep")
    m = load_manifest("t_sleep")
    m["wall_seconds"] = 2
    save_manifest(m)
    t0 = time.time()
    rid, state, _ = run_tool("t_sleep", [])
    took = time.time() - t0
    row = run_row(rid)
    check("wall clock kills the run", state == "TIMEOUT"
          and row[0] == "TIMEOUT" and took < 15)
    check("timeout audited", len(events_of("RUN_TIMEOUT", rid)) == 1)

    mktool("t_fail", TOOL_FAIL)
    cmd_enroll("t_fail")
    rid, state, _ = run_tool("t_fail", [])
    row = run_row(rid)
    check("nonzero exit -> FAILED with code",
          state == "FAILED" and row[1] == 3)
    # --- door lifecycle (bug 6) --------------------------------------------
    dead = _p("runs", rid, "door", "door.sock")
    check("door.sock removed on normal shutdown", not os.path.lexists(dead))
    def _node(p):
        s = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM); s.bind(p); s.close()
    _node(dead)                          # listener died without stop()
    live_rid = "rlivedecoy001"
    os.makedirs(_p("runs", live_rid, "door"), exist_ok=True)
    live_sock = _p("runs", live_rid, "door", "door.sock")
    _node(live_sock)
    def _decoy(c):
        c.execute("INSERT INTO runs(rid,tool,argv,state,room,image,created,"
                  "invoker) VALUES(?,?,?,?,?,?,?,?)",
                  (live_rid, "t_fail", "[]", "RUNNING", "x", "-", now(), me))
    audited("RUN_CREATE", live_rid, "selftest live decoy", fn=_decoy)
    n = _sweep_dead_doors()
    check("sweep removes dead door of terminal run", n == 1 and not os.path.lexists(dead))
    check("sweep never touches a RUNNING door", os.path.lexists(live_sock))
    _node(dead)                          # restart on the same run dir
    b = DoorBroker(rid, load_manifest("t_fail"), dead); b.start(); b.stop(); b.join(timeout=5)
    check("broker rebinds over stale door, cleans up", not os.path.lexists(dead))
    def _close_decoy(c):
        c.execute("UPDATE runs SET state='DONE', ended=? WHERE rid=?", (now(), live_rid))
    audited("RUN_DONE", live_rid, "selftest decoy closed", fn=_close_decoy)
    os.unlink(live_sock)

    # --- room seam (argv only; no podman here) ---------------------------------
    argv = _room_cmd("tb-x", "/kit/t", "/r/s", "/r/d", "img:1", ["a", "b"])
    check("room argv: network none", "--network=none" in argv)
    check("room argv: keep-id (rmtree fix baked in)",
          "--userns=keep-id" in argv)
    check("room argv: caps dropped, no-new-privs",
          "ALL" in argv and "no-new-privileges" in argv)
    check("room argv: tool ro, scratch rw, door mounted",
          "/kit/t:/tool:ro" in argv and "/r/s:/scratch:rw" in argv
          and "/r/d:/door:rw" in argv)
    check("room argv: entry + args intact",
          argv[-4:] == ["python3", "/tool/main.py", "a", "b"])
    check("room argv: podman by absolute path (T2 rev 3)",
          argv[0] == "/usr/bin/podman")
    _saved = {k: os.environ.get(k) for k in ("CONTAINERS_CONF_OVERRIDE", "XDG_CONFIG_HOME")}
    os.environ["CONTAINERS_CONF_OVERRIDE"] = "/tmp/evil.conf"
    os.environ["XDG_CONFIG_HOME"] = "/tmp/evil-xdg"
    _re = _room_env()
    for _k, _v in _saved.items():
        if _v is None:
            os.environ.pop(_k, None)
        else:
            os.environ[_k] = _v
    check("room env: root-owned containers.conf only, no override, fixed XDG config",
          _re.get("CONTAINERS_CONF") == "/usr/share/containers/containers.conf"
          and "CONTAINERS_CONF_OVERRIDE" not in _re
          and _re.get("XDG_CONFIG_HOME") == "/opt/treebeard/shared/toolbox/podman-config")

    # --- hygiene ---------------------------------------------------------------
    staging = [f for f in os.listdir(_p("staging"))
               if f != "probe-tool"]
    check("staging provably empty", staging == [])
    card_txt = open(_p("status", "toolbox.txt"), encoding="ascii").read()
    check("card flags UNSAFE room loud", "UNSAFE" in card_txt)
    conn = db()
    try:
        dangling = conn.execute(
            "SELECT COUNT(*) FROM door_calls WHERE state='SPENT'"
        ).fetchone()[0]
        no_terminal = conn.execute(
            "SELECT COUNT(*) FROM runs WHERE state IN ('PENDING','RUNNING')"
        ).fetchone()[0]
    finally:
        conn.close()
    check("no dangling SPENT door calls", dangling == 0)
    check("no runs left PENDING or RUNNING", no_terminal == 0)
    conn = db()
    try:
        rids = [r[0] for r in conn.execute("SELECT rid FROM runs")]
        covered = True
        for r in rids:
            kinds = set(k[0] for k in conn.execute(
                "SELECT kind FROM events WHERE rid=?", (r,)))
            if "RUN_CREATE" not in kinds or not any(
                    k.startswith("RUN_") and k != "RUN_CREATE"
                    and k != "RUN_START" for k in kinds):
                covered = False
    finally:
        conn.close()
    check("audit trail covers every run", covered)

    print("SELFTEST %d/%d %s" % (passed[0], passed[0] + failed[0],
                                 "PASS" if failed[0] == 0 else "FAIL"))
    shutil.rmtree(tmp, ignore_errors=True)
    if failed[0]:
        sys.exit(1)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

USAGE = """usage: toolbox <command>
  init                                   create the data tree
  enroll <tool>                          deny-all manifest for a shelf tool
  allow <tool> <skill> [k=v ...]         grant a door service
      router:     tiers=free,cheap max_calls=4 max_tokens=1200
      state:      cards=meeseeks,anvil max_calls=8
      web_search: max_calls=3 max_results=5
  revoke <tool> <skill>                  remove a grant
  invokers <tool> <u1,u2>                who may run the tool
  run <tool> [args...]                   run inside the room
  status [tool]                          card, runs, door calls, log tails
  skills                                 list knowledge skill docs
  cards                                  list the state-card registry
  probe                                  prove the real wall + door + keep-id
  selftest                               full offline proof (no podman)
"""


def main(argv):
    if len(argv) < 2:
        sys.stderr.write(USAGE)
        return 2
    cmd = argv[1]
    try:
        if cmd == "init":
            cmd_init()
        elif cmd == "enroll" and len(argv) == 3:
            cmd_enroll(argv[2])
        elif cmd == "allow" and len(argv) >= 4:
            cmd_allow(argv[2], argv[3], argv[4:])
        elif cmd == "revoke" and len(argv) == 4:
            cmd_revoke(argv[2], argv[3])
        elif cmd == "invokers" and len(argv) == 4:
            cmd_invokers(argv[2], argv[3])
        elif cmd == "run" and len(argv) >= 3:
            cmd_run(argv[2], argv[3:])
        elif cmd == "status":
            cmd_status(argv[2] if len(argv) > 2 else None)
        elif cmd == "skills":
            cmd_skills()
        elif cmd == "cards":
            cmd_cards()
        elif cmd == "probe":
            cmd_probe()
        elif cmd == "selftest":
            cmd_selftest()
        else:
            sys.stderr.write(USAGE)
            return 2
    except (ValueError, KeyError, FileNotFoundError, FileExistsError,
            PermissionError, RuntimeError) as e:
        sys.stderr.write("toolbox: %s\n" % e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
